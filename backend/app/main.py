from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import uuid
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from zoneinfo import ZoneInfo

import httpx
from fastapi import FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

try:
    from .aliyun_asr import aliyun_asr_enabled, is_aliyun_asr_provider, recognize_speech_with_aliyun
    from .aliyun_ocr import (
        aliyun_ocr_enabled,
        is_aliyun_ocr_provider,
        recognize_answer_sheet_with_aliyun,
        recognize_material_text_with_aliyun,
    )
    from .textbook_knowledge import TEXTBOOK_CHAR_META, YEAR_ONE_LESSONS
except ImportError:  # pragma: no cover - supports direct script execution in local debugging.
    from aliyun_asr import aliyun_asr_enabled, is_aliyun_asr_provider, recognize_speech_with_aliyun
    from aliyun_ocr import (
        aliyun_ocr_enabled,
        is_aliyun_ocr_provider,
        recognize_answer_sheet_with_aliyun,
        recognize_material_text_with_aliyun,
    )
    from textbook_knowledge import TEXTBOOK_CHAR_META, YEAR_ONE_LESSONS


APP_NAME = "AI 生字词智能过关小助手"
DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./data/app.db")
AI_API_BASE = os.getenv("AI_API_BASE", "").rstrip("/")
AI_API_KEY = os.getenv("AI_API_KEY", "")
AI_MODEL = os.getenv("AI_MODEL", "qwen-plus")
AI_PROVIDER = os.getenv("AI_PROVIDER", "openai_compatible")
OCR_PROVIDER = os.getenv("OCR_PROVIDER", "mock")
OCR_API_URL = os.getenv("OCR_API_URL", "").rstrip("/")
OCR_API_KEY = os.getenv("OCR_API_KEY", "")
ASR_PROVIDER = os.getenv("ASR_PROVIDER", "mock")
ASR_API_URL = os.getenv("ASR_API_URL", "").rstrip("/")
ASR_API_KEY = os.getenv("ASR_API_KEY", "")


def now_iso() -> str:
    return datetime.utcnow().isoformat(timespec="seconds") + "Z"


def format_beijing_time(value: Optional[str]) -> str:
    if not value:
        return ""
    try:
        normalized = value.replace("Z", "+00:00")
        dt = datetime.fromisoformat(normalized)
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return value


def new_id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


def json_dumps(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def json_loads(value: Optional[str], default: Any = None) -> Any:
    if not value:
        return default
    try:
        return json.loads(value)
    except json.JSONDecodeError:
        return default


class Database:
    def __init__(self, url: str):
        self.url = url
        self.is_postgres = url.startswith("postgresql://") or url.startswith("postgres://")
        if not self.is_postgres and url.startswith("sqlite:///"):
            db_path = Path(url.replace("sqlite:///", "", 1))
            if not db_path.is_absolute():
                db_path = Path.cwd() / db_path
            db_path.parent.mkdir(parents=True, exist_ok=True)
            self.sqlite_path = str(db_path)
        else:
            self.sqlite_path = ""

    @contextmanager
    def connect(self):
        if self.is_postgres:
            try:
                import psycopg
                from psycopg.rows import dict_row
            except ModuleNotFoundError as exc:
                raise RuntimeError("PostgreSQL 模式需要安装 psycopg[binary]") from exc
            conn = psycopg.connect(self.url, row_factory=dict_row)
        else:
            conn = sqlite3.connect(self.sqlite_path)
            conn.row_factory = sqlite3.Row
        try:
            yield conn
            conn.commit()
        finally:
            conn.close()

    def sql(self, statement: str) -> str:
        if not self.is_postgres:
            return statement
        return statement.replace("%", "%%").replace("?", "%s")

    def execute(self, statement: str, params: Iterable[Any] = ()) -> None:
        with self.connect() as conn:
            conn.execute(self.sql(statement), tuple(params))

    def many(self, statement: str, rows: Iterable[Iterable[Any]]) -> None:
        with self.connect() as conn:
            if self.is_postgres:
                with conn.cursor() as cur:
                    cur.executemany(self.sql(statement), [tuple(row) for row in rows])
            else:
                conn.executemany(statement, [tuple(row) for row in rows])

    def query(self, statement: str, params: Iterable[Any] = ()) -> List[Dict[str, Any]]:
        with self.connect() as conn:
            cur = conn.execute(self.sql(statement), tuple(params))
            return [dict(row) for row in cur.fetchall()]

    def one(self, statement: str, params: Iterable[Any] = ()) -> Optional[Dict[str, Any]]:
        with self.connect() as conn:
            cur = conn.execute(self.sql(statement), tuple(params))
            row = cur.fetchone()
            return dict(row) if row else None


db = Database(DATABASE_URL)


SCHEMA = [
    """
    CREATE TABLE IF NOT EXISTS users (
      id TEXT PRIMARY KEY,
      name TEXT NOT NULL,
      role TEXT NOT NULL,
      class_id TEXT,
      created_at TEXT NOT NULL
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS classes (
      id TEXT PRIMARY KEY,
      name TEXT NOT NULL,
      grade TEXT NOT NULL,
      teacher_id TEXT NOT NULL,
      created_at TEXT NOT NULL
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS lessons (
      id TEXT PRIMARY KEY,
      grade TEXT NOT NULL,
      volume TEXT NOT NULL,
      unit_no INTEGER NOT NULL,
      title TEXT NOT NULL,
      content TEXT NOT NULL,
      chars_json TEXT NOT NULL
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS learning_packs (
      id TEXT PRIMARY KEY,
      lesson_id TEXT,
      creator_id TEXT NOT NULL,
      title TEXT NOT NULL,
      grade TEXT NOT NULL,
      volume TEXT NOT NULL,
      unit_name TEXT NOT NULL,
      source_type TEXT NOT NULL,
      status TEXT NOT NULL,
      content_json TEXT NOT NULL,
      ai_generation_id TEXT,
      created_at TEXT NOT NULL,
      confirmed_at TEXT
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS dictation_tasks (
      id TEXT PRIMARY KEY,
      class_id TEXT NOT NULL,
      teacher_id TEXT NOT NULL,
      learning_pack_id TEXT NOT NULL,
      title TEXT NOT NULL,
      mode TEXT NOT NULL,
      status TEXT NOT NULL,
      settings_json TEXT NOT NULL,
      created_at TEXT NOT NULL,
      deadline TEXT,
      published_at TEXT
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS dictation_items (
      id TEXT PRIMARY KEY,
      task_id TEXT NOT NULL,
      item_type TEXT NOT NULL,
      prompt_text TEXT NOT NULL,
      answer TEXT NOT NULL,
      answer_meta_json TEXT NOT NULL,
      audio_text TEXT NOT NULL,
      order_no INTEGER NOT NULL,
      difficulty INTEGER NOT NULL
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS submissions (
      id TEXT PRIMARY KEY,
      task_id TEXT NOT NULL,
      student_id TEXT NOT NULL,
      status TEXT NOT NULL,
      total_count INTEGER NOT NULL,
      correct_count INTEGER NOT NULL,
      wrong_count INTEGER NOT NULL,
      suspected_count INTEGER NOT NULL,
      pending_review_count INTEGER NOT NULL,
      created_at TEXT NOT NULL,
      submitted_at TEXT
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS submission_answers (
      id TEXT PRIMARY KEY,
      submission_id TEXT NOT NULL,
      item_id TEXT NOT NULL,
      raw_answer TEXT NOT NULL,
      normalized_answer TEXT NOT NULL,
      recognized_text TEXT,
      confidence REAL,
      result TEXT NOT NULL,
      mistake_type TEXT,
      feedback_json TEXT NOT NULL,
      created_at TEXT NOT NULL,
      reviewed_by TEXT,
      reviewed_at TEXT
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS mistake_records (
      id TEXT PRIMARY KEY,
      student_id TEXT NOT NULL,
      char_or_word TEXT NOT NULL,
      char_id TEXT,
      mistake_type TEXT NOT NULL,
      wrong_count INTEGER NOT NULL,
      source_item_ids_json TEXT NOT NULL,
      first_wrong_at TEXT NOT NULL,
      last_wrong_at TEXT NOT NULL,
      status TEXT NOT NULL,
      next_review_at TEXT
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS ai_generation_records (
      id TEXT PRIMARY KEY,
      scene TEXT NOT NULL,
      model_provider TEXT NOT NULL,
      model_name TEXT NOT NULL,
      prompt_version TEXT NOT NULL,
      input_hash TEXT NOT NULL,
      output_json TEXT NOT NULL,
      latency_ms INTEGER NOT NULL,
      status TEXT NOT NULL,
      created_by TEXT NOT NULL,
      created_at TEXT NOT NULL
    )
    """,
    """
    CREATE TABLE IF NOT EXISTS pronunciation_records (
      id TEXT PRIMARY KEY,
      student_id TEXT NOT NULL,
      item_id TEXT,
      target_text TEXT NOT NULL,
      recognized_text TEXT NOT NULL,
      expected_pinyin TEXT NOT NULL,
      score INTEGER NOT NULL,
      mastery TEXT NOT NULL,
      issues_json TEXT NOT NULL,
      correction_json TEXT NOT NULL,
      provider TEXT NOT NULL,
      created_at TEXT NOT NULL
    )
    """,
]


KNOWN_CHARS: Dict[str, Dict[str, Any]] = {
    "晴": {
        "pinyin": "qing2",
        "display_pinyin": "qíng",
        "radical": "日",
        "structure": "左右结构",
        "stroke_count": 12,
        "stroke_order": ["竖", "横折", "横", "横", "横", "横", "竖", "横", "竖", "横折钩", "横", "横"],
        "words": ["晴天", "晴朗", "放晴"],
        "simple_sentence": "今天是晴天。",
        "confusing_chars": [{"char": "睛", "reason": "睛是目字旁，和眼睛有关；晴是日字旁，和天气有关。"}],
        "common_mistakes": ["容易把日字旁写成目字旁。"],
    },
    "睛": {
        "pinyin": "jing1",
        "display_pinyin": "jīng",
        "radical": "目",
        "structure": "左右结构",
        "stroke_count": 13,
        "stroke_order": ["竖", "横折", "横", "横", "横", "横", "横", "竖", "横", "竖", "横折钩", "横", "横"],
        "words": ["眼睛", "目不转睛"],
        "simple_sentence": "我有一双明亮的眼睛。",
        "confusing_chars": [{"char": "晴", "reason": "睛是目字旁，表示和眼睛有关。"}],
        "common_mistakes": ["右边是青，左边是目字旁。"],
    },
    "情": {
        "pinyin": "qing2",
        "display_pinyin": "qíng",
        "radical": "忄",
        "structure": "左右结构",
        "stroke_count": 11,
        "stroke_order": ["点", "点", "竖", "横", "横", "竖", "横", "竖", "横折钩", "横", "横"],
        "words": ["心情", "友情", "事情"],
        "simple_sentence": "我的心情很好。",
        "confusing_chars": [{"char": "晴", "reason": "情是竖心旁，常和心情、感情有关。"}],
        "common_mistakes": ["左边是竖心旁，不是日字旁。"],
    },
    "请": {
        "pinyin": "qing3",
        "display_pinyin": "qǐng",
        "radical": "讠",
        "structure": "左右结构",
        "stroke_count": 10,
        "stroke_order": ["点", "横折提", "横", "横", "竖", "横", "竖", "横折钩", "横", "横"],
        "words": ["请问", "请坐", "邀请"],
        "simple_sentence": "请你坐下。",
        "confusing_chars": [{"char": "情", "reason": "请是言字旁，和说话、请求有关。"}],
        "common_mistakes": ["左边是言字旁，不是竖心旁。"],
    },
    "清": {
        "pinyin": "qing1",
        "display_pinyin": "qīng",
        "radical": "氵",
        "structure": "左右结构",
        "stroke_count": 11,
        "stroke_order": ["点", "点", "提", "横", "横", "竖", "横", "竖", "横折钩", "横", "横"],
        "words": ["清水", "清早", "清楚"],
        "simple_sentence": "小河的水很清。",
        "confusing_chars": [{"char": "请", "reason": "清是三点水，常和水或干净有关。"}],
        "common_mistakes": ["左边是三点水。"],
    },
    "已": {
        "pinyin": "yi3",
        "display_pinyin": "yǐ",
        "radical": "己",
        "structure": "独体字",
        "stroke_count": 3,
        "stroke_order": ["横折", "横", "竖弯钩"],
        "words": ["已经", "已知"],
        "simple_sentence": "我已经写完作业了。",
        "confusing_chars": [{"char": "己", "reason": "已字上面开口，己字下面不开口。"}],
        "common_mistakes": ["容易和己、巳混淆。"],
    },
    "己": {
        "pinyin": "ji3",
        "display_pinyin": "jǐ",
        "radical": "己",
        "structure": "独体字",
        "stroke_count": 3,
        "stroke_order": ["横折", "横", "竖弯钩"],
        "words": ["自己", "知己"],
        "simple_sentence": "自己的事情自己做。",
        "confusing_chars": [{"char": "已", "reason": "己字下面不开口，已字上面开口。"}],
        "common_mistakes": ["容易和已混淆。"],
    },
    "生": {
        "pinyin": "sheng1",
        "display_pinyin": "shēng",
        "radical": "生",
        "structure": "独体字",
        "stroke_count": 5,
        "stroke_order": ["撇", "横", "横", "竖", "横"],
        "words": ["学生", "生日", "生活"],
        "simple_sentence": "今天是我的生日。",
        "confusing_chars": [],
        "common_mistakes": ["中间一竖要写正。"],
    },
    "字": {
        "pinyin": "zi4",
        "display_pinyin": "zì",
        "radical": "宀",
        "structure": "上下结构",
        "stroke_count": 6,
        "stroke_order": ["点", "点", "横钩", "横撇", "竖钩", "横"],
        "words": ["生字", "写字", "汉字"],
        "simple_sentence": "我会认真写字。",
        "confusing_chars": [],
        "common_mistakes": ["宝盖头不要写得太宽。"],
    },
    "词": {
        "pinyin": "ci2",
        "display_pinyin": "cí",
        "radical": "讠",
        "structure": "左右结构",
        "stroke_count": 7,
        "stroke_order": ["点", "横折提", "横折钩", "横", "竖", "横折", "横"],
        "words": ["词语", "组词"],
        "simple_sentence": "我们学习新的词语。",
        "confusing_chars": [],
        "common_mistakes": ["左边是言字旁。"],
    },
}


KNOWN_CHARS.update({char: meta for char, meta in TEXTBOOK_CHAR_META.items() if char not in KNOWN_CHARS})

DEMO_LESSON_CHARS = ["晴", "睛", "情", "请", "清", "已", "己", "生", "字", "词"]
DEMO_TEACHER_ID = "teacher_demo"
DEMO_CLASS_ID = "class_1_1"
DEMO_PACK_ID = "pack_demo_qing"
DEMO_TASK_ID = "task_demo_qing"
DEMO_GENERATION_ID = "ai_demo_qing"
TEXTBOOK_LESSON_CHARS = list(
    dict.fromkeys(char for lesson in YEAR_ONE_LESSONS for char in lesson.get("chars", []))
)
KNOWLEDGE_BASE_FALLBACK_CHARS = list(dict.fromkeys(DEMO_LESSON_CHARS + TEXTBOOK_LESSON_CHARS))


def clean_storage_text(value: Any) -> str:
    return str(value or "").strip()


def init_db() -> None:
    for statement in SCHEMA:
        db.execute(statement)
    seed_demo_data()


def seed_demo_data(force: bool = False, full_demo: bool = False) -> Dict[str, Any]:
    if force:
        for table in [
            "submission_answers",
            "submissions",
            "mistake_records",
            "dictation_items",
            "dictation_tasks",
            "learning_packs",
            "ai_generation_records",
            "pronunciation_records",
            "lessons",
            "users",
            "classes",
        ]:
            db.execute(f"DELETE FROM {table}")
    existing = db.one("SELECT id FROM users LIMIT 1")
    if existing:
        ensure_material_lessons_seeded()
        return {"seeded": False, "class_id": None, "task_id": None}

    created_at = now_iso()
    teacher_id = DEMO_TEACHER_ID
    class_id = DEMO_CLASS_ID
    students = [
        ("student_chen", "陈小雨"),
        ("student_lin", "林一凡"),
        ("student_wang", "王可欣"),
        ("student_zhou", "周明明"),
        ("student_li", "李安安"),
        ("student_zhao", "赵乐乐"),
    ]
    users = [(teacher_id, "张老师", "teacher", None, created_at)]
    users.extend((sid, name, "student", class_id, created_at) for sid, name in students)
    db.many(
        "INSERT INTO users (id,name,role,class_id,created_at) VALUES (?,?,?,?,?)",
        users,
    )
    db.execute(
        "INSERT INTO classes (id,name,grade,teacher_id,created_at) VALUES (?,?,?,?,?)",
        (class_id, "一年级 1 班", "一年级", teacher_id, created_at),
    )
    ensure_material_lessons_seeded()
    summary = {"seeded": True, "class_id": class_id, "task_id": None}
    if full_demo:
        summary.update(seed_full_demo_flow(created_at))
    return summary


def ensure_teacher_exists(teacher_id: str) -> Dict[str, Any]:
    normalized_id = clean_storage_text(teacher_id) or DEMO_TEACHER_ID
    teacher = db.one("SELECT * FROM users WHERE id=? AND role='teacher'", (normalized_id,))
    if teacher:
        return teacher
    db.execute(
        "INSERT INTO users (id,name,role,class_id,created_at) VALUES (?,?,?,?,?)",
        (normalized_id, "张老师" if normalized_id == DEMO_TEACHER_ID else "默认教师", "teacher", None, now_iso()),
    )
    return db.one("SELECT * FROM users WHERE id=? AND role='teacher'", (normalized_id,)) or {
        "id": normalized_id,
        "name": "默认教师",
        "role": "teacher",
        "class_id": None,
        "created_at": now_iso(),
    }


def seed_full_demo_flow(created_at: str) -> Dict[str, Any]:
    content = "晴天里，小朋友看着清清的小河，心情很好。请大家认真学习生字和词语。"
    lesson_chars = [build_char_payload(ch) for ch in DEMO_LESSON_CHARS]
    payload = build_learning_pack_payload("一年级", "下册", "第一单元", "识字练习：天气和心情", content, lesson_chars)
    dictation_items = payload["dictation_items"][:8]
    selected_answers = [item["answer"] for item in dictation_items]

    db.execute(
        """
        INSERT INTO ai_generation_records
        (id,scene,model_provider,model_name,prompt_version,input_hash,output_json,latency_ms,status,created_by,created_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            DEMO_GENERATION_ID,
            "learning_pack",
            "demo_seed",
            AI_MODEL,
            "learning_pack_v1",
            hashlib.sha256(json_dumps(payload).encode("utf-8")).hexdigest(),
            json_dumps(payload),
            0,
            "fallback",
            DEMO_TEACHER_ID,
            created_at,
        ),
    )
    db.execute(
        """
        INSERT INTO learning_packs
        (id,lesson_id,creator_id,title,grade,volume,unit_name,source_type,status,content_json,ai_generation_id,created_at,confirmed_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            DEMO_PACK_ID,
            "lesson_demo_qing",
            DEMO_TEACHER_ID,
            "识字练习：天气和心情",
            "一年级",
            "下册",
            "第一单元",
            "textbook",
            "confirmed",
            json_dumps(payload),
            DEMO_GENERATION_ID,
            created_at,
            created_at,
        ),
    )

    published_at = now_iso()
    db.execute(
        """
        INSERT INTO dictation_tasks
        (id,class_id,teacher_id,learning_pack_id,title,mode,status,settings_json,created_at,deadline,published_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            DEMO_TASK_ID,
            DEMO_CLASS_ID,
            DEMO_TEACHER_ID,
            DEMO_PACK_ID,
            "识字练习：天气和心情 AI 听写任务",
            "classroom",
            "published",
            json_dumps(
                {
                    "speed": 0.9,
                    "repeat": 2,
                    "interval_seconds": 10,
                    "question_type": "word",
                    "selected_answers": selected_answers,
                }
            ),
            created_at,
            (datetime.utcnow() + timedelta(days=7)).isoformat(timespec="seconds") + "Z",
            published_at,
        ),
    )

    item_rows = []
    item_ids: List[str] = []
    for index, item in enumerate(dictation_items, start=1):
        answer = clean_storage_text(item.get("answer"))
        char = clean_storage_text(item.get("char")) or (answer[0] if answer else "")
        item_id = f"item_demo_qing_{index:02d}"
        item_ids.append(item_id)
        item_rows.append(
            (
                item_id,
                DEMO_TASK_ID,
                "word",
                clean_storage_text(item.get("prompt_text")) or f"请写：{answer}",
                answer,
                json_dumps({"char": char, "difficulty": item.get("difficulty", 1), "original_answer": answer, "question_type": "word"}),
                clean_dictation_audio_text(clean_storage_text(item.get("audio_text")), answer),
                index,
                int(item.get("difficulty", 1)),
            )
        )
    db.many(
        """
        INSERT INTO dictation_items
        (id,task_id,item_type,prompt_text,answer,answer_meta_json,audio_text,order_no,difficulty)
        VALUES (?,?,?,?,?,?,?,?,?)
        """,
        item_rows,
    )

    answer_by_item = {row[0]: row[4] for row in item_rows}

    def wrong_answer(answer: str) -> str:
        replacements = {"眼睛": "眼晴", "心情": "心晴", "请问": "清问", "清水": "请水", "自己": "自已"}
        return replacements.get(answer, inject_demo_error(answer))

    def seed_submission(student_id: str, suffix: str, wrong_indexes: Iterable[int] = (), pending_indexes: Iterable[int] = (), answered_count: Optional[int] = None) -> None:
        submission_id = f"sub_demo_{suffix}"
        total_count = len(item_ids)
        limit = answered_count or total_count
        db.execute(
            """
            INSERT INTO submissions
            (id,task_id,student_id,status,total_count,correct_count,wrong_count,suspected_count,pending_review_count,created_at,submitted_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?)
            """,
            (submission_id, DEMO_TASK_ID, student_id, "in_progress", total_count, 0, 0, 0, 0, created_at, None),
        )
        wrong_set = set(wrong_indexes)
        pending_set = set(pending_indexes)
        for order_no, item_id in enumerate(item_ids[:limit], start=1):
            answer = answer_by_item[item_id]
            if order_no in pending_set:
                record_answer(submission_id, item_id, answer, 0.62)
            elif order_no in wrong_set:
                record_answer(submission_id, item_id, wrong_answer(answer), 0.91)
            else:
                record_answer(submission_id, item_id, answer, 0.96)
        status = "in_progress" if limit < total_count else "submitted"
        submitted_at = None if status == "in_progress" else now_iso()
        db.execute("UPDATE submissions SET status=?, submitted_at=? WHERE id=?", (status, submitted_at, submission_id))

    seed_submission("student_zhou", "zhou", wrong_indexes=())
    seed_submission("student_li", "li", wrong_indexes=(2, 5))
    seed_submission("student_lin", "lin", wrong_indexes=(3,), pending_indexes=(4,))
    seed_submission("student_wang", "wang", wrong_indexes=(1,), answered_count=4)
    seed_submission("student_zhao", "zhao", wrong_indexes=())

    pronunciation_cases = [
        ("student_zhou", item_ids[0], selected_answers[0], selected_answers[0], 0.95, "mock"),
        ("student_li", item_ids[1], selected_answers[1], wrong_answer(selected_answers[1]), 0.72, "mock"),
        ("student_lin", item_ids[2], selected_answers[2], selected_answers[2], 0.88, "mock"),
    ]
    for student_id, item_id, target_text, recognized_text, confidence, provider in pronunciation_cases:
        result = evaluate_pronunciation_result(target_text, recognized_text, confidence)
        db.execute(
            """
            INSERT INTO pronunciation_records
            (id,student_id,item_id,target_text,recognized_text,expected_pinyin,score,mastery,issues_json,correction_json,provider,created_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                new_id("pron"),
                student_id,
                item_id,
                target_text,
                result["recognized_text"],
                result["expected_pinyin"],
                result["score"],
                result["mastery"],
                json_dumps(result["issues"]),
                json_dumps(result["correction"]),
                provider,
                now_iso(),
            ),
        )

    return {
        "learning_pack_id": DEMO_PACK_ID,
        "task_id": DEMO_TASK_ID,
        "student_count": 6,
        "dictation_item_count": len(item_ids),
    }


def build_char_payload(ch: str) -> Dict[str, Any]:
    base = KNOWN_CHARS.get(ch, {})
    return {
        "char": ch,
        "pinyin": base.get("pinyin", ""),
        "display_pinyin": base.get("display_pinyin", ""),
        "radical": base.get("radical", "待确认"),
        "structure": base.get("structure", "待确认"),
        "stroke_count": base.get("stroke_count"),
        "stroke_order": base.get("stroke_order", []),
        "words": base.get("words", [f"{ch}字"]),
        "simple_sentence": base.get("simple_sentence", ""),
        "confusing_chars": base.get("confusing_chars", []),
        "common_mistakes": base.get("common_mistakes", ["请教师确认易错点。"]),
        "dictation_level": "basic",
        "needs_teacher_review": base.get("needs_teacher_review", ch not in KNOWN_CHARS),
    }


def upsert_lesson(
    lesson_id: str,
    grade: str,
    volume: str,
    unit_no: int,
    title: str,
    content: str,
    chars: Iterable[str],
) -> None:
    chars_json = json_dumps([build_char_payload(ch) for ch in chars])
    existing = db.one("SELECT id FROM lessons WHERE id=?", (lesson_id,))
    if existing:
        db.execute(
            """
            UPDATE lessons
            SET grade=?, volume=?, unit_no=?, title=?, content=?, chars_json=?
            WHERE id=?
            """,
            (grade, volume, unit_no, title, content, chars_json, lesson_id),
        )
        return
    db.execute(
        """
        INSERT INTO lessons (id,grade,volume,unit_no,title,content,chars_json)
        VALUES (?,?,?,?,?,?,?)
        """,
        (lesson_id, grade, volume, unit_no, title, content, chars_json),
    )


def upsert_textbook_lessons() -> None:
    for item in YEAR_ONE_LESSONS:
        upsert_lesson(
            item["id"],
            item["grade"],
            item["volume"],
            int(item["unit_no"]),
            item["title"],
            item["content"],
            item["chars"],
        )


def upsert_demo_lesson() -> None:
    upsert_lesson(
        "lesson_demo_qing",
        "一年级",
        "下册",
        1,
        "识字练习：天气和心情",
        "晴天里，小朋友看着清清的小河，心情很好。请大家认真学习生字和词语。",
        DEMO_LESSON_CHARS,
    )


def ensure_material_lessons_seeded() -> None:
    lesson_count = db.one("SELECT COUNT(*) AS count FROM lessons")
    current_count = int((lesson_count or {}).get("count") or 0)
    if current_count < len(YEAR_ONE_LESSONS):
        upsert_textbook_lessons()
    upsert_demo_lesson()


def extract_chars(text: str) -> List[str]:
    seen: List[str] = []
    for ch in re.findall(r"[\u4e00-\u9fff]", text):
        if ch in KNOWN_CHARS and ch not in seen:
            seen.append(ch)
    for ch in KNOWLEDGE_BASE_FALLBACK_CHARS:
        if len(seen) >= 10:
            break
        if ch not in seen:
            seen.append(ch)
    return seen[:12]


def build_learning_pack_payload(
    grade: str,
    volume: str,
    unit: str,
    title: str,
    content: str,
    lesson_chars: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, Any]:
    characters = lesson_chars or [build_char_payload(ch) for ch in extract_chars(f"{title}{content}")]
    dictation_items: List[Dict[str, Any]] = []
    for index, item in enumerate(characters[:10], start=1):
        words = item.get("words") or [item["char"]]
        answer = words[0]
        dictation_items.append(
            {
                "type": "word",
                "answer": answer,
                "prompt_text": f"请写词语：{answer}",
                "audio_text": answer,
                "difficulty": 1 if index <= 5 else 2,
                "char": item["char"],
            }
        )
    return {
        "lesson": {"grade": grade, "volume": volume, "unit": unit, "title": title},
        "characters": characters,
        "dictation_items": dictation_items,
        "review_notes": ["演示模式已生成可直接发布的生字词学习包，正式使用前建议教师确认。"],
    }


async def call_ai_learning_pack(request: "GenerateLearningPackRequest", lesson_chars: Optional[List[Dict[str, Any]]]) -> Tuple[Dict[str, Any], str]:
    fallback = build_learning_pack_payload(
        request.grade,
        request.volume,
        request.unit,
        request.title,
        request.content or "",
        lesson_chars,
    )
    if not AI_API_BASE or not AI_API_KEY:
        return fallback, "fallback"

    prompt = f"""
你是一名小学一二年级语文教师助手。请根据输入内容生成低年级适用的生字词学习包。
要求：
1. 只输出 JSON，不要输出 Markdown。
2. 内容必须适合小学一二年级。
3. 每个生字给出拼音、部首、结构、笔画数、2-4 个组词、一个简单例句、形近字辨析和易错提醒。
4. 与给定标准生字冲突时，将 needs_teacher_review 设为 true。

年级：{request.grade}
册次：{request.volume}
单元：{request.unit}
标题：{request.title}
课文或教师输入：{request.content}
标准生字参考：{json_dumps(lesson_chars or [])}

JSON 结构：
{{
  "lesson": {{"grade": "", "volume": "", "unit": "", "title": ""}},
  "characters": [
    {{
      "char": "",
      "pinyin": "qing2",
      "display_pinyin": "qíng",
      "radical": "",
      "structure": "",
      "stroke_count": 0,
      "words": [],
      "simple_sentence": "",
      "confusing_chars": [{{"char": "", "reason": ""}}],
      "common_mistakes": [],
      "dictation_level": "basic",
      "needs_teacher_review": false
    }}
  ],
  "dictation_items": [
    {{"type": "word", "answer": "", "prompt_text": "", "audio_text": "", "difficulty": 1, "char": ""}}
  ],
  "review_notes": []
}}
""".strip()
    url = AI_API_BASE
    if not url.endswith("/chat/completions"):
        url = f"{url}/chat/completions"
    payload = {
        "model": AI_MODEL,
        "messages": [
            {"role": "system", "content": "你是严谨的小学语文 AI 教学工具，只输出合法 JSON。"},
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.2,
        "response_format": {"type": "json_object"},
    }
    try:
        async with httpx.AsyncClient(timeout=45) as client:
            response = await client.post(
                url,
                headers={"Authorization": f"Bearer {AI_API_KEY}", "Content-Type": "application/json"},
                json=payload,
            )
            response.raise_for_status()
        raw = response.json()["choices"][0]["message"]["content"]
        generated = json.loads(raw)
        if validate_learning_pack(generated):
            return generated, "success"
    except Exception:
        return fallback, "fallback"
    return fallback, "fallback"


def validate_learning_pack(payload: Dict[str, Any]) -> bool:
    if not isinstance(payload.get("characters"), list) or not payload["characters"]:
        return False
    for item in payload["characters"]:
        if not all(key in item for key in ["char", "display_pinyin", "radical", "structure", "words"]):
            return False
    if not isinstance(payload.get("dictation_items"), list) or not payload["dictation_items"]:
        return False
    return True


def normalize_answer(value: str) -> str:
    value = value.strip()
    value = re.sub(r"[\s，。,.、；;：:！!？?\"'“”‘’（）()【】\[\]]+", "", value)
    return value


def flatten_confusing_chars() -> Dict[str, List[str]]:
    pairs: Dict[str, List[str]] = {}
    for ch, data in KNOWN_CHARS.items():
        for item in data.get("confusing_chars", []):
            pairs.setdefault(ch, []).append(item["char"])
            pairs.setdefault(item["char"], []).append(ch)
    return pairs


CONFUSING_MAP = flatten_confusing_chars()


def pinyin_of(value: str) -> str:
    if len(value) == 1 and value in KNOWN_CHARS:
        return KNOWN_CHARS[value]["pinyin"]
    return ""


def classify_mistake(answer: str, standard: str) -> str:
    if answer == standard:
        return "unknown"
    if not answer:
        return "unknown"
    if len(answer) == 1 and len(standard) == 1:
        if answer in CONFUSING_MAP.get(standard, []):
            return "similar_shape_confusion"
        if pinyin_of(answer) and pinyin_of(answer) == pinyin_of(standard):
            return "homophone_confusion"
    if len(standard) > 1 and len(answer) == len(standard):
        for got, expected in zip(answer, standard):
            reason = classify_mistake(got, expected)
            if reason != "unknown":
                return reason
    return "glyph_error"


def feedback_for(raw_answer: str, standard: str, result: str, mistake_type: Optional[str]) -> Dict[str, Any]:
    if result == "correct":
        return {"message": "答对了，继续保持。", "correct_answer": standard, "tips": []}
    tips = [f"正确答案是：{standard}。"]
    main_char = standard[0] if standard else ""
    if main_char in KNOWN_CHARS:
        data = KNOWN_CHARS[main_char]
        tips.append(f"{main_char} 的拼音是 {data['display_pinyin']}，部首是 {data['radical']}。")
        if data.get("common_mistakes"):
            tips.append(data["common_mistakes"][0])
    if mistake_type == "similar_shape_confusion":
        tips.append("这个错误属于形近字混淆，可以重点看偏旁和意思。")
    if mistake_type == "homophone_confusion":
        tips.append("这个错误属于同音字混淆，可以放进句子里辨一辨。")
    return {"message": "这题需要再练一次。", "correct_answer": standard, "tips": tips}


def inject_demo_error(standard: str) -> str:
    for index, char in enumerate(standard):
        confusing = CONFUSING_MAP.get(char, [])
        if confusing:
            return standard[:index] + confusing[0] + standard[index + 1 :]
    if standard:
        return standard[:-1] if len(standard) > 1 else ""
    return ""


def word_display_pinyin(value: str) -> str:
    result: List[str] = []
    for char in value:
        data = KNOWN_CHARS.get(char)
        result.append(data["display_pinyin"] if data else char)
    return " ".join(result)


def clean_dictation_audio_text(value: str, fallback: str = "") -> str:
    text = (value or fallback or "").strip()
    for prefix in ("请写词语：", "请写词语:", "请写：", "请写:", "请写这个字：", "请写这个字:"):
        if text.startswith(prefix):
            return text[len(prefix) :].strip()
    return text


def choice_options_for(char: str) -> List[str]:
    options = [char] if char else []
    for confusing in CONFUSING_MAP.get(char, []):
        if confusing not in options:
            options.append(confusing)
    for fallback in KNOWLEDGE_BASE_FALLBACK_CHARS:
        if len(options) >= 4:
            break
        if fallback != char and fallback not in options:
            options.append(fallback)
    return options[:4]


async def call_generic_ocr(file: UploadFile, expected_answers: List[str]) -> Tuple[List[str], float, str]:
    filename = (file.filename or "").lower()
    content = await file.read()
    if is_aliyun_ocr_provider(OCR_PROVIDER) and aliyun_ocr_enabled(OCR_PROVIDER):
        try:
            result = await asyncio.to_thread(recognize_answer_sheet_with_aliyun, content, expected_answers)
            if result.answers:
                return result.answers, result.confidence, result.provider
        except Exception:
            pass
    if OCR_API_URL and OCR_API_KEY:
        try:
            async with httpx.AsyncClient(timeout=45) as client:
                response = await client.post(
                    OCR_API_URL,
                    headers={"Authorization": f"Bearer {OCR_API_KEY}"},
                    data={"expected_answers": json_dumps(expected_answers)},
                    files={"image": (file.filename or "answer-sheet.jpg", content, file.content_type or "application/octet-stream")},
                )
                response.raise_for_status()
            payload = response.json()
            raw_answers = payload.get("answers") or payload.get("recognized_answers") or payload.get("texts") or []
            if isinstance(raw_answers, str):
                raw_answers = re.split(r"[\n,，、;；]+", raw_answers)
            answers = [str(item).strip() for item in raw_answers if str(item).strip()]
            confidence = float(payload.get("confidence", 0.9))
            if answers:
                return answers, confidence, OCR_PROVIDER or "cloud"
        except Exception:
            pass

    answers = list(expected_answers)
    if "wrong" in filename or "cuo" in filename or "错" in filename:
        if len(answers) >= 2:
            answers[1] = inject_demo_error(answers[1])
        elif answers:
            answers[0] = inject_demo_error(answers[0])
    elif "blank" in filename or "low" in filename:
        return ["" for _ in expected_answers], 0.62, "mock_low_confidence"
    return answers, 0.92, "mock"


def extract_text_from_ocr_payload(payload: Dict[str, Any]) -> str:
    for key in ("extracted_text", "recognized_text", "text", "content"):
        value = payload.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    for key in ("answers", "recognized_answers", "texts", "lines"):
        value = payload.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
        if isinstance(value, list):
            lines = [str(item).strip() for item in value if str(item).strip()]
            if lines:
                return "\n".join(lines)
    return ""


async def call_material_text_ocr(file: UploadFile) -> Tuple[str, float, str]:
    content = await file.read()
    if is_aliyun_ocr_provider(OCR_PROVIDER) and aliyun_ocr_enabled(OCR_PROVIDER):
        try:
            result = await asyncio.to_thread(recognize_material_text_with_aliyun, content)
            if result.raw_text:
                return result.raw_text, result.confidence, result.provider
        except Exception:
            pass
    if OCR_API_URL:
        try:
            headers = {"Authorization": f"Bearer {OCR_API_KEY}"} if OCR_API_KEY else {}
            async with httpx.AsyncClient(timeout=45) as client:
                response = await client.post(
                    OCR_API_URL,
                    headers=headers,
                    data={"scene": "textbook_extract", "return_text": "true"},
                    files={"image": (file.filename or "textbook-page.jpg", content, file.content_type or "application/octet-stream")},
                )
                response.raise_for_status()
            payload = response.json()
            text = extract_text_from_ocr_payload(payload)
            if text:
                return text, float(payload.get("confidence", 0.88)), OCR_PROVIDER or "cloud"
        except Exception:
            pass

    try:
        decoded = content.decode("utf-8").strip()
        if decoded and len(decoded) <= 3000:
            return decoded, 0.96, "local_text_file"
    except UnicodeDecodeError:
        pass

    return "晴天里，小朋友看着清清的小河，心情很好。请大家认真学习生字和词语。", 0.78, "mock_textbook_ocr"


async def call_generic_asr(file: UploadFile, target_text: str) -> Tuple[str, float, str]:
    filename = (file.filename or "").lower()
    content = await file.read()
    if is_aliyun_asr_provider(ASR_PROVIDER) and aliyun_asr_enabled(ASR_PROVIDER):
        try:
            result = await asyncio.to_thread(
                recognize_speech_with_aliyun,
                content,
                file.filename or "pronunciation.wav",
                file.content_type or "",
            )
            if result.recognized_text:
                return result.recognized_text, result.confidence, result.provider
        except Exception:
            pass

    if ASR_API_URL and ASR_API_KEY:
        try:
            async with httpx.AsyncClient(timeout=45) as client:
                response = await client.post(
                    ASR_API_URL,
                    headers={"Authorization": f"Bearer {ASR_API_KEY}"},
                    data={"target_text": target_text},
                    files={"audio": (file.filename or "pronunciation.webm", content, file.content_type or "application/octet-stream")},
                )
                response.raise_for_status()
            payload = response.json()
            recognized = str(payload.get("recognized_text") or payload.get("text") or payload.get("transcript") or "").strip()
            confidence = float(payload.get("confidence", 0.9))
            if recognized:
                return recognized, confidence, ASR_PROVIDER or "cloud"
        except Exception:
            pass

    if "wrong" in filename or "cuo" in filename or "错" in filename:
        return inject_demo_error(target_text) or target_text[:-1], 0.68, "mock"
    return target_text, 0.94, "mock"


def evaluate_pronunciation_result(target_text: str, recognized_text: str, confidence: float) -> Dict[str, Any]:
    target = normalize_answer(target_text)
    recognized = normalize_answer(recognized_text)
    exact = target == recognized
    if exact:
        score = max(85, min(98, round(confidence * 100)))
        issues: List[str] = []
    else:
        diff_count = abs(len(target) - len(recognized))
        for got, expected in zip(recognized, target):
            if got != expected:
                diff_count += 1
        score = max(45, min(82, round(confidence * 75) - diff_count * 4))
        issues = ["识别结果与目标词不一致，需要重点听清声母、韵母和声调。"]
        for got, expected in zip(recognized, target):
            if got != expected and expected in KNOWN_CHARS:
                data = KNOWN_CHARS[expected]
                issues.append(f"“{expected}”应读作 {data['display_pinyin']}。")
                break
    mastery = "passed" if score >= 85 else "needs_practice" if score >= 70 else "weak"
    correction = [
        f"目标词：{target_text}，标准拼音：{word_display_pinyin(target_text)}。",
        "建议先慢读一遍，再连起来读完整词语。",
    ]
    if not exact:
        correction.append("可以跟读 AI 标准音，再重新录音。")
    return {
        "recognized_text": recognized_text,
        "expected_pinyin": word_display_pinyin(target_text),
        "score": score,
        "mastery": mastery,
        "issues": issues,
        "correction": correction,
    }


def upsert_mistake(student_id: str, standard: str, item_id: str, mistake_type: str) -> None:
    existing = db.one(
        "SELECT * FROM mistake_records WHERE student_id=? AND char_or_word=? AND mistake_type=?",
        (student_id, standard, mistake_type),
    )
    now = now_iso()
    if existing:
        wrong_count = int(existing["wrong_count"]) + 1
        sources = json_loads(existing["source_item_ids_json"], [])
        if item_id not in sources:
            sources.append(item_id)
        next_review = (datetime.utcnow() + timedelta(days=1 if wrong_count <= 2 else 3)).isoformat(timespec="seconds") + "Z"
        db.execute(
            """
            UPDATE mistake_records
            SET wrong_count=?, source_item_ids_json=?, last_wrong_at=?, status=?, next_review_at=?
            WHERE id=?
            """,
            (wrong_count, json_dumps(sources), now, "uncorrected", next_review, existing["id"]),
        )
    else:
        db.execute(
            """
            INSERT INTO mistake_records
            (id,student_id,char_or_word,char_id,mistake_type,wrong_count,source_item_ids_json,first_wrong_at,last_wrong_at,status,next_review_at)
            VALUES (?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                new_id("mistake"),
                student_id,
                standard,
                standard[0] if standard else None,
                mistake_type,
                1,
                json_dumps([item_id]),
                now,
                now,
                "uncorrected",
                (datetime.utcnow() + timedelta(days=1)).isoformat(timespec="seconds") + "Z",
            ),
        )


class GenerateLearningPackRequest(BaseModel):
    grade: str = "一年级"
    volume: str = "下册"
    unit: str = "第一单元"
    title: str = "识字练习：天气和心情"
    content: str = ""
    lesson_id: Optional[str] = None
    teacher_id: str = "teacher_demo"


class ConfirmLearningPackRequest(BaseModel):
    teacher_id: str = "teacher_demo"


class CreateClassRequest(BaseModel):
    name: str
    grade: str = "一年级"
    teacher_id: str = "teacher_demo"


class ImportStudentsRequest(BaseModel):
    names: List[str]


class CreateTaskRequest(BaseModel):
    class_id: str = "class_1_1"
    teacher_id: str = "teacher_demo"
    learning_pack_id: str
    title: str
    mode: str = "classroom"
    question_type: str = "word"
    selected_answers: Optional[List[str]] = None
    custom_items: Optional[List[Dict[str, Any]]] = None
    deadline: Optional[str] = None
    settings: Dict[str, Any] = Field(default_factory=lambda: {"speed": 0.9, "repeat": 1, "interval_seconds": 8})


class CreateSubmissionRequest(BaseModel):
    student_id: str


class SubmitAnswerRequest(BaseModel):
    item_id: str
    raw_answer: str
    confidence: Optional[float] = 1.0


class ReviewRequest(BaseModel):
    student_id: str


class UpdateMistakeStatusRequest(BaseModel):
    student_id: str
    status: str = Field(pattern="^(uncorrected|corrected|reviewing|passed|repeated)$")


class ReviewAnswerRequest(BaseModel):
    result: str = Field(pattern="^(correct|wrong)$")
    corrected_answer: Optional[str] = None
    teacher_id: str = "teacher_demo"


class BatchReviewRequest(BaseModel):
    answer_ids: List[str]
    result: str = Field(pattern="^(correct|wrong)$")
    corrected_answer: Optional[str] = None
    teacher_id: str = "teacher_demo"


app = FastAPI(title=APP_NAME, version="0.1.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=os.getenv("CORS_ALLOW_ORIGINS", "*").split(","),
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def startup() -> None:
    init_db()


@app.get("/api/health")
def health() -> Dict[str, Any]:
    ocr_cloud_ready = (OCR_API_URL and OCR_API_KEY) or aliyun_ocr_enabled(OCR_PROVIDER)
    asr_cloud_ready = (ASR_API_URL and ASR_API_KEY) or aliyun_asr_enabled(ASR_PROVIDER)
    return {
        "ok": True,
        "name": APP_NAME,
        "database": "postgresql" if db.is_postgres else "sqlite",
        "ai_mode": "cloud" if AI_API_BASE and AI_API_KEY else "fallback",
        "ocr_mode": "cloud" if ocr_cloud_ready else "mock",
        "ocr_provider": OCR_PROVIDER,
        "asr_mode": "cloud" if asr_cloud_ready else "mock",
        "asr_provider": ASR_PROVIDER,
        "time": now_iso(),
    }


@app.post("/api/demo/reset")
def reset_demo() -> Dict[str, Any]:
    summary = seed_demo_data(force=True, full_demo=True)
    return {"ok": True, **summary}


@app.get("/api/classes")
def list_classes() -> List[Dict[str, Any]]:
    return db.query("SELECT * FROM classes ORDER BY created_at DESC")


@app.post("/api/classes")
def create_class(request: CreateClassRequest) -> Dict[str, Any]:
    teacher = ensure_teacher_exists(request.teacher_id)
    class_name = clean_storage_text(request.name)
    grade = clean_storage_text(request.grade)
    if not class_name:
        raise HTTPException(status_code=400, detail="班级名称不能为空")
    if not grade:
        raise HTTPException(status_code=400, detail="年级不能为空")
    existing = db.one(
        "SELECT * FROM classes WHERE teacher_id=? AND grade=? AND name=?",
        (teacher["id"], grade, class_name),
    )
    if existing:
        raise HTTPException(status_code=409, detail=f"{grade}已存在班级：{class_name}")
    class_id = new_id("class")
    db.execute(
        "INSERT INTO classes (id,name,grade,teacher_id,created_at) VALUES (?,?,?,?,?)",
        (class_id, class_name, grade, teacher["id"], now_iso()),
    )
    row = db.one("SELECT * FROM classes WHERE id=?", (class_id,))
    return row or {"id": class_id, "name": class_name, "grade": grade, "teacher_id": teacher["id"]}


@app.get("/api/classes/{class_id}/students")
def list_students(class_id: str) -> List[Dict[str, Any]]:
    return db.query("SELECT id,name,role,class_id,created_at FROM users WHERE role='student' AND class_id=? ORDER BY name", (class_id,))


@app.post("/api/classes/{class_id}/students/import")
def import_students(class_id: str, request: ImportStudentsRequest) -> Dict[str, Any]:
    class_row = db.one("SELECT id FROM classes WHERE id=?", (class_id,))
    if not class_row:
        raise HTTPException(status_code=404, detail="班级不存在")
    names = []
    seen = set()
    for name in request.names:
        cleaned = clean_storage_text(name)
        if cleaned and cleaned not in seen:
            seen.add(cleaned)
            names.append(cleaned)
    if not names:
        raise HTTPException(status_code=400, detail="学生名单不能为空")
    now = now_iso()
    rows = [(new_id("student"), name, "student", class_id, now) for name in names]
    db.many("INSERT INTO users (id,name,role,class_id,created_at) VALUES (?,?,?,?,?)", rows)
    return {"imported_count": len(rows), "students": list_students(class_id)}


@app.get("/api/materials/lessons")
def list_lessons(grade: Optional[str] = None) -> List[Dict[str, Any]]:
    ensure_material_lessons_seeded()
    order_sql = "grade, CASE volume WHEN '上册' THEN 1 WHEN '下册' THEN 2 ELSE 3 END, unit_no, CASE WHEN id LIKE 'pep_%' THEN 0 ELSE 1 END, id, title"
    if grade:
        rows = db.query(f"SELECT * FROM lessons WHERE grade=? ORDER BY {order_sql}", (grade,))
    else:
        rows = db.query(f"SELECT * FROM lessons ORDER BY {order_sql}")
    for row in rows:
        row["chars"] = json_loads(row.pop("chars_json"), [])
    return rows


@app.post("/api/materials/extract-image")
async def extract_material_image(image: UploadFile = File(...)) -> Dict[str, Any]:
    text, confidence, provider = await call_material_text_ocr(image)
    return {
        "provider": provider,
        "confidence": confidence,
        "extracted_text": text,
        "filename": image.filename,
    }


@app.post("/api/learning-packs/generate")
async def generate_learning_pack(request: GenerateLearningPackRequest) -> Dict[str, Any]:
    lesson_chars = None
    lesson = None
    if request.lesson_id:
        lesson = db.one("SELECT * FROM lessons WHERE id=?", (request.lesson_id,))
        if not lesson:
            raise HTTPException(status_code=404, detail="课文不存在")
        lesson_chars = json_loads(lesson["chars_json"], [])
    started = datetime.utcnow()
    payload, status = await call_ai_learning_pack(request, lesson_chars)
    latency_ms = int((datetime.utcnow() - started).total_seconds() * 1000)
    generation_id = new_id("ai")
    input_hash = hashlib.sha256(json_dumps(request.model_dump()).encode("utf-8")).hexdigest()
    db.execute(
        """
        INSERT INTO ai_generation_records
        (id,scene,model_provider,model_name,prompt_version,input_hash,output_json,latency_ms,status,created_by,created_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            generation_id,
            "learning_pack",
            AI_PROVIDER,
            AI_MODEL,
            "learning_pack_v1",
            input_hash,
            json_dumps(payload),
            latency_ms,
            status,
            request.teacher_id,
            now_iso(),
        ),
    )
    pack_id = new_id("pack")
    lesson_meta = payload.get("lesson", {})
    db.execute(
        """
        INSERT INTO learning_packs
        (id,lesson_id,creator_id,title,grade,volume,unit_name,source_type,status,content_json,ai_generation_id,created_at,confirmed_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            pack_id,
            request.lesson_id,
            request.teacher_id,
            request.title,
            lesson_meta.get("grade", request.grade),
            lesson_meta.get("volume", request.volume),
            lesson_meta.get("unit", request.unit),
            "textbook" if lesson else "pasted_text",
            "pending_review",
            json_dumps(payload),
            generation_id,
            now_iso(),
            None,
        ),
    )
    return {"id": pack_id, "status": "pending_review", "ai_status": status, "content": payload}


@app.get("/api/learning-packs/{pack_id}")
def get_learning_pack(pack_id: str) -> Dict[str, Any]:
    row = db.one("SELECT * FROM learning_packs WHERE id=?", (pack_id,))
    if not row:
        raise HTTPException(status_code=404, detail="学习包不存在")
    row["content"] = json_loads(row.pop("content_json"), {})
    return row


@app.put("/api/learning-packs/{pack_id}")
def update_learning_pack(pack_id: str, content: Dict[str, Any]) -> Dict[str, Any]:
    row = db.one("SELECT id FROM learning_packs WHERE id=?", (pack_id,))
    if not row:
        raise HTTPException(status_code=404, detail="学习包不存在")
    db.execute("UPDATE learning_packs SET content_json=?, status=? WHERE id=?", (json_dumps(content), "pending_review", pack_id))
    return get_learning_pack(pack_id)


@app.post("/api/learning-packs/{pack_id}/confirm")
def confirm_learning_pack(pack_id: str, request: ConfirmLearningPackRequest) -> Dict[str, Any]:
    row = db.one("SELECT id FROM learning_packs WHERE id=? AND creator_id=?", (pack_id, request.teacher_id))
    if not row:
        raise HTTPException(status_code=404, detail="学习包不存在或无权限")
    db.execute("UPDATE learning_packs SET status=?, confirmed_at=? WHERE id=?", ("confirmed", now_iso(), pack_id))
    return get_learning_pack(pack_id)


@app.post("/api/dictation-tasks")
def create_dictation_task(request: CreateTaskRequest) -> Dict[str, Any]:
    pack = get_learning_pack(request.learning_pack_id)
    question_type = request.question_type if request.question_type in {"word", "char", "pinyin", "pinyin_to_word", "choice"} else "word"
    items = pack["content"].get("dictation_items", [])
    custom_items = [item for item in (request.custom_items or []) if normalize_answer(item.get("answer", ""))]
    using_custom_items = bool(custom_items)
    if using_custom_items:
        items = custom_items
    elif request.selected_answers:
        selected = {normalize_answer(answer) for answer in request.selected_answers}
        filtered_items = []
        for item in items:
            original_answer = normalize_answer(item.get("answer", ""))
            char_answer = normalize_answer(item.get("char") or (original_answer[:1] if original_answer else ""))
            if question_type in {"char", "choice"}:
                matched = original_answer in selected or char_answer in selected
            else:
                matched = original_answer in selected
            if matched:
                filtered_items.append(item)
        items = filtered_items
    if not items:
        raise HTTPException(status_code=400, detail="听写范围不能为空")
    task_id = new_id("task")
    settings = dict(request.settings or {})
    settings.setdefault("speed", 0.9)
    settings["repeat"] = max(1, int(settings.get("repeat", 1) or 1))
    settings["interval_seconds"] = max(3, int(settings.get("interval_seconds", 8) or 8))
    settings["question_type"] = question_type
    if request.selected_answers is not None:
        settings["selected_answers"] = request.selected_answers
    if request.deadline:
        settings["deadline"] = request.deadline
    db.execute(
        """
        INSERT INTO dictation_tasks
        (id,class_id,teacher_id,learning_pack_id,title,mode,status,settings_json,created_at,deadline,published_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            task_id,
            request.class_id,
            request.teacher_id,
            request.learning_pack_id,
            request.title,
            request.mode,
            "draft",
            json_dumps(settings),
            now_iso(),
            request.deadline,
            None,
        ),
    )
    rows = []
    seen_task_answers = set()
    for index, item in enumerate(items, start=1):
        item_type = item.get("type") if item.get("type") in {"word", "char", "pinyin", "pinyin_to_word", "choice"} else question_type
        if using_custom_items:
            answer = clean_storage_text(item.get("answer"))
            original_answer = clean_storage_text(item.get("source_answer") or item.get("original_answer") or answer)
            char = clean_storage_text(item.get("char") or (answer[0] if answer else ""))
            prompt = clean_storage_text(item.get("prompt_text")) or f"请写：{answer}"
            audio_text = clean_dictation_audio_text(clean_storage_text(item.get("audio_text")), answer)
        else:
            original_answer = item.get("answer", "")
            char = item.get("char") or (original_answer[0] if original_answer else "")
            if question_type == "char":
                answer = char or original_answer
                prompt = f"请写这个字：{answer}"
                audio_text = answer
            elif question_type == "pinyin":
                answer = word_display_pinyin(original_answer)
                prompt = f"请写拼音：{original_answer}"
                audio_text = original_answer
            elif question_type == "pinyin_to_word":
                answer = original_answer
                prompt = f"看拼音写词语：{word_display_pinyin(original_answer)}"
                audio_text = word_display_pinyin(original_answer)
            elif question_type == "choice":
                answer = char or original_answer[:1]
                prompt = "听音选字：请从选项中选择正确的字"
                audio_text = answer
            else:
                answer = original_answer
                prompt = item.get("prompt_text") or f"请写：{answer}"
                audio_text = clean_dictation_audio_text(item.get("audio_text", ""), answer)
        if item_type in {"char", "choice"}:
            task_answer_key = normalize_answer(answer)
            if task_answer_key in seen_task_answers:
                continue
            seen_task_answers.add(task_answer_key)
        meta = {"char": char, "difficulty": item.get("difficulty", 1), "original_answer": original_answer, "question_type": item_type}
        if item_type == "pinyin_to_word":
            meta["display_pinyin"] = word_display_pinyin(original_answer)
        if item_type == "choice":
            meta["options"] = choice_options_for(answer)
        rows.append(
            (
                new_id("item"),
                task_id,
                item_type,
                prompt,
                answer,
                json_dumps(meta),
                audio_text,
                len(rows) + 1,
                int(item.get("difficulty", 1)),
            )
        )
    if not rows:
        raise HTTPException(status_code=400, detail="听写范围不能为空")
    db.many(
        """
        INSERT INTO dictation_items
        (id,task_id,item_type,prompt_text,answer,answer_meta_json,audio_text,order_no,difficulty)
        VALUES (?,?,?,?,?,?,?,?,?)
        """,
        rows,
    )
    return get_dictation_task(task_id)


@app.post("/api/dictation-tasks/{task_id}/publish")
def publish_dictation_task(task_id: str) -> Dict[str, Any]:
    row = db.one("SELECT id FROM dictation_tasks WHERE id=?", (task_id,))
    if not row:
        raise HTTPException(status_code=404, detail="任务不存在")
    db.execute("UPDATE dictation_tasks SET status=?, published_at=? WHERE id=?", ("published", now_iso(), task_id))
    return get_dictation_task(task_id)


@app.get("/api/dictation-tasks/{task_id}")
def get_dictation_task(task_id: str) -> Dict[str, Any]:
    row = db.one("SELECT * FROM dictation_tasks WHERE id=?", (task_id,))
    if not row:
        raise HTTPException(status_code=404, detail="任务不存在")
    row["settings"] = json_loads(row.pop("settings_json"), {})
    items = db.query("SELECT * FROM dictation_items WHERE task_id=? ORDER BY order_no", (task_id,))
    for item in items:
        item["answer_meta"] = json_loads(item.pop("answer_meta_json"), {})
    row["items"] = items
    return row


@app.get("/api/classes/{class_id}/dictation-tasks")
def list_class_tasks(class_id: str) -> List[Dict[str, Any]]:
    rows = db.query("SELECT * FROM dictation_tasks WHERE class_id=? ORDER BY created_at DESC", (class_id,))
    for row in rows:
        row["settings"] = json_loads(row.pop("settings_json"), {})
    return rows


def preferred_submission(task_id: str, student_id: str) -> Optional[Dict[str, Any]]:
    return db.one(
        """
        SELECT *
        FROM submissions
        WHERE task_id=? AND student_id=?
        ORDER BY
          CASE
            WHEN status IN ('submitted','reviewed') THEN 0
            WHEN status='in_progress' THEN 1
            ELSE 2
          END,
          COALESCE(submitted_at, created_at) DESC,
          created_at DESC
        LIMIT 1
        """,
        (task_id, student_id),
    )


@app.get("/api/student/tasks")
def list_student_tasks(student_id: str = Query(...)) -> List[Dict[str, Any]]:
    student = db.one("SELECT * FROM users WHERE id=? AND role='student'", (student_id,))
    if not student:
        raise HTTPException(status_code=404, detail="学生不存在")
    rows = db.query(
        "SELECT * FROM dictation_tasks WHERE class_id=? AND status='published' ORDER BY published_at DESC",
        (student["class_id"],),
    )
    for row in rows:
        row["settings"] = json_loads(row.pop("settings_json"), {})
        submission = preferred_submission(row["id"], student_id)
        row["submission"] = submission
    return rows


@app.get("/api/student/tasks/{task_id}")
def get_student_task(task_id: str, student_id: str = Query(...)) -> Dict[str, Any]:
    task = get_dictation_task(task_id)
    student = db.one("SELECT * FROM users WHERE id=? AND role='student'", (student_id,))
    if not student or student["class_id"] != task["class_id"]:
        raise HTTPException(status_code=403, detail="无权访问该任务")
    return task


@app.post("/api/student/tasks/{task_id}/submissions")
def create_submission(task_id: str, request: CreateSubmissionRequest) -> Dict[str, Any]:
    task = get_student_task(task_id, request.student_id)
    existing = preferred_submission(task_id, request.student_id)
    if existing:
        return get_submission_result(existing["id"])
    submission_id = new_id("sub")
    total_count = len(task["items"])
    db.execute(
        """
        INSERT INTO submissions
        (id,task_id,student_id,status,total_count,correct_count,wrong_count,suspected_count,pending_review_count,created_at,submitted_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?)
        """,
        (submission_id, task_id, request.student_id, "in_progress", total_count, 0, 0, 0, 0, now_iso(), None),
    )
    return get_submission_result(submission_id)


def record_answer(submission_id: str, item_id: str, raw_answer: str, confidence: Optional[float] = 1.0) -> Dict[str, Any]:
    submission = db.one("SELECT * FROM submissions WHERE id=?", (submission_id,))
    if not submission:
        raise HTTPException(status_code=404, detail="提交记录不存在")
    item = db.one("SELECT * FROM dictation_items WHERE id=?", (item_id,))
    if not item:
        raise HTTPException(status_code=404, detail="题目不存在")
    if item["task_id"] != submission["task_id"]:
        raise HTTPException(status_code=400, detail="题目不属于该任务")
    existing = db.one("SELECT id FROM submission_answers WHERE submission_id=? AND item_id=?", (submission_id, item_id))
    if existing:
        db.execute("DELETE FROM submission_answers WHERE id=?", (existing["id"],))

    normalized = normalize_answer(raw_answer)
    standard = normalize_answer(item["answer"])
    if confidence is not None and confidence < 0.7:
        result = "pending_review"
        mistake_type = "unknown"
    elif normalized == standard:
        result = "correct"
        mistake_type = None
    else:
        result = "wrong"
        mistake_type = classify_mistake(normalized, standard)
        upsert_mistake(submission["student_id"], standard, item["id"], mistake_type)

    feedback = feedback_for(normalized, standard, result, mistake_type)
    answer_id = new_id("ans")
    db.execute(
        """
        INSERT INTO submission_answers
        (id,submission_id,item_id,raw_answer,normalized_answer,recognized_text,confidence,result,mistake_type,feedback_json,created_at,reviewed_by,reviewed_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
                answer_id,
                submission_id,
                item["id"],
                raw_answer,
                normalized,
                normalized,
                confidence,
                result,
                mistake_type,
                json_dumps(feedback),
            now_iso(),
            None,
            None,
        ),
    )
    refresh_submission_counts(submission_id)
    return get_answer(answer_id)


@app.post("/api/submissions/{submission_id}/answers")
def submit_answer(submission_id: str, request: SubmitAnswerRequest) -> Dict[str, Any]:
    return record_answer(submission_id, request.item_id, request.raw_answer, request.confidence)


@app.post("/api/submissions/{submission_id}/answers/image-sheet")
async def submit_answer_sheet(
    submission_id: str,
    image: UploadFile = File(...),
    recognized_texts: Optional[str] = Form(None),
) -> Dict[str, Any]:
    submission = db.one("SELECT * FROM submissions WHERE id=?", (submission_id,))
    if not submission:
        raise HTTPException(status_code=404, detail="提交记录不存在")
    items = db.query("SELECT * FROM dictation_items WHERE task_id=? ORDER BY order_no", (submission["task_id"],))
    if not items:
        raise HTTPException(status_code=400, detail="任务没有题目")

    expected_answers = [item["answer"] for item in items]
    if recognized_texts:
        recognized_answers = [part.strip() for part in re.split(r"[\n,，、;；]+", recognized_texts) if part.strip()]
        confidence = 0.96
        provider = "manual_review_text"
    else:
        recognized_answers, confidence, provider = await call_generic_ocr(image, expected_answers)
    if len(recognized_answers) < len(items):
        recognized_answers = recognized_answers + [""] * (len(items) - len(recognized_answers))

    answers: List[Dict[str, Any]] = []
    for item, raw_answer in zip(items, recognized_answers[: len(items)]):
        answers.append(record_answer(submission_id, item["id"], raw_answer, confidence))
    refreshed = get_submission_result(submission_id)
    return {
        "provider": provider,
        "confidence": confidence,
        "recognized_answers": recognized_answers[: len(items)],
        "answers": answers,
        "submission": refreshed,
    }


@app.post("/api/submissions/{submission_id}/answers/image")
async def submit_answer_image(
    submission_id: str,
    item_id: str = Form(...),
    image: UploadFile = File(...),
    recognized_text: Optional[str] = Form(None),
) -> Dict[str, Any]:
    item = db.one("SELECT * FROM dictation_items WHERE id=?", (item_id,))
    if not item:
        raise HTTPException(status_code=404, detail="题目不存在")
    if recognized_text:
        raw_answer = recognized_text
        confidence = 0.96
        provider = "manual_review_text"
    else:
        answers, confidence, provider = await call_generic_ocr(image, [item["answer"]])
        raw_answer = answers[0] if answers else ""
    answer = record_answer(submission_id, item_id, raw_answer, confidence)
    return {"provider": provider, "confidence": confidence, "recognized_text": raw_answer, "answer": answer}


@app.post("/api/pronunciation/evaluate")
async def evaluate_pronunciation(
    student_id: str = Form(...),
    target_text: str = Form(...),
    item_id: Optional[str] = Form(None),
    audio: UploadFile = File(...),
) -> Dict[str, Any]:
    student = db.one("SELECT id FROM users WHERE id=? AND role='student'", (student_id,))
    if not student:
        raise HTTPException(status_code=404, detail="学生不存在")
    recognized_text, confidence, provider = await call_generic_asr(audio, target_text)
    result = evaluate_pronunciation_result(target_text, recognized_text, confidence)
    record_id = new_id("pron")
    db.execute(
        """
        INSERT INTO pronunciation_records
        (id,student_id,item_id,target_text,recognized_text,expected_pinyin,score,mastery,issues_json,correction_json,provider,created_at)
        VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            record_id,
            student_id,
            item_id,
            target_text,
            result["recognized_text"],
            result["expected_pinyin"],
            result["score"],
            result["mastery"],
            json_dumps(result["issues"]),
            json_dumps(result["correction"]),
            provider,
            now_iso(),
        ),
    )
    return {"id": record_id, "provider": provider, "confidence": confidence, **result}


def get_answer(answer_id: str) -> Dict[str, Any]:
    row = db.one("SELECT * FROM submission_answers WHERE id=?", (answer_id,))
    if not row:
        raise HTTPException(status_code=404, detail="作答不存在")
    row["feedback"] = json_loads(row.pop("feedback_json"), {})
    item = db.one("SELECT id,item_type,prompt_text,answer,audio_text,order_no,answer_meta_json FROM dictation_items WHERE id=?", (row["item_id"],))
    if item:
        item["answer_meta"] = json_loads(item.pop("answer_meta_json"), {})
    row["item"] = item
    return row


def refresh_submission_counts(submission_id: str) -> None:
    answers = db.query("SELECT result FROM submission_answers WHERE submission_id=?", (submission_id,))
    correct = sum(1 for item in answers if item["result"] == "correct")
    wrong = sum(1 for item in answers if item["result"] == "wrong")
    suspected = sum(1 for item in answers if item["result"] == "suspected")
    pending = sum(1 for item in answers if item["result"] == "pending_review")
    db.execute(
        """
        UPDATE submissions
        SET correct_count=?, wrong_count=?, suspected_count=?, pending_review_count=?
        WHERE id=?
        """,
        (correct, wrong, suspected, pending, submission_id),
    )


def review_submission_answer(answer_id: str, request: ReviewAnswerRequest) -> Dict[str, Any]:
    answer = db.one(
        """
        SELECT sa.*, s.student_id, s.status AS submission_status
        FROM submission_answers sa
        JOIN submissions s ON sa.submission_id = s.id
        WHERE sa.id=?
        """,
        (answer_id,),
    )
    if not answer:
        raise HTTPException(status_code=404, detail="作答不存在")
    item = db.one("SELECT * FROM dictation_items WHERE id=?", (answer["item_id"],))
    if not item:
        raise HTTPException(status_code=404, detail="题目不存在")

    standard = normalize_answer(item["answer"])
    corrected_raw = request.corrected_answer if request.corrected_answer is not None else answer.get("raw_answer", "")
    corrected = normalize_answer(corrected_raw)
    if request.result == "correct":
        result = "correct"
        normalized = standard
        mistake_type = None
        feedback = feedback_for(standard, standard, result, mistake_type)
        raw_answer = corrected_raw or item["answer"]
    else:
        result = "wrong"
        normalized = corrected
        mistake_type = classify_mistake(normalized, standard)
        feedback = feedback_for(normalized, standard, result, mistake_type)
        raw_answer = corrected_raw
        upsert_mistake(answer["student_id"], standard, item["id"], mistake_type)

    reviewed_at = now_iso()
    db.execute(
        """
        UPDATE submission_answers
        SET raw_answer=?, normalized_answer=?, recognized_text=?, result=?, mistake_type=?, feedback_json=?, reviewed_by=?, reviewed_at=?
        WHERE id=?
        """,
        (
            raw_answer,
            normalized,
            normalized,
            result,
            mistake_type,
            json_dumps(feedback),
            request.teacher_id,
            reviewed_at,
            answer_id,
        ),
    )
    refresh_submission_counts(answer["submission_id"])
    counts = db.one("SELECT pending_review_count,status FROM submissions WHERE id=?", (answer["submission_id"],))
    if counts and int(counts["pending_review_count"]) == 0 and counts["status"] == "submitted":
        db.execute("UPDATE submissions SET status=? WHERE id=?", ("reviewed", answer["submission_id"]))
    return get_answer(answer_id)


@app.get("/api/review/pending-answers")
def list_pending_review_answers(class_id: str = Query("class_1_1"), task_id: Optional[str] = None) -> List[Dict[str, Any]]:
    params: List[Any] = [class_id]
    task_clause = ""
    if task_id:
        task_clause = "AND dt.id=?"
        params.append(task_id)
    rows = db.query(
        f"""
        SELECT
          sa.id, sa.submission_id, sa.item_id, sa.raw_answer, sa.normalized_answer, sa.recognized_text,
          sa.confidence, sa.result, sa.mistake_type, sa.created_at,
          s.student_id, u.name AS student_name,
          dt.id AS task_id, dt.title AS task_title,
          di.prompt_text, di.answer, di.audio_text, di.order_no, di.item_type
        FROM submission_answers sa
        JOIN submissions s ON sa.submission_id = s.id
        JOIN users u ON s.student_id = u.id
        JOIN dictation_items di ON sa.item_id = di.id
        JOIN dictation_tasks dt ON di.task_id = dt.id
        WHERE dt.class_id=?
          {task_clause}
          AND (sa.result IN ('pending_review','suspected') OR COALESCE(sa.confidence, 1) < 0.7)
        ORDER BY sa.created_at DESC
        """,
        params,
    )
    for row in rows:
        row["student"] = {"id": row["student_id"], "name": row["student_name"]}
        row["task"] = {"id": row["task_id"], "title": row["task_title"]}
        row["item"] = {
            "id": row["item_id"],
            "prompt_text": row["prompt_text"],
            "answer": row["answer"],
            "audio_text": row["audio_text"],
            "order_no": row["order_no"],
            "item_type": row["item_type"],
        }
    return rows


@app.post("/api/review/answers/{answer_id}")
def review_answer(answer_id: str, request: ReviewAnswerRequest) -> Dict[str, Any]:
    return review_submission_answer(answer_id, request)


@app.post("/api/review/answers/batch")
def batch_review_answers(request: BatchReviewRequest) -> Dict[str, Any]:
    reviewed = []
    for answer_id in request.answer_ids:
        reviewed.append(
            review_submission_answer(
                answer_id,
                ReviewAnswerRequest(result=request.result, corrected_answer=request.corrected_answer, teacher_id=request.teacher_id),
            )
        )
    return {"reviewed_count": len(reviewed), "answers": reviewed}


@app.post("/api/submissions/{submission_id}/finish")
def finish_submission(submission_id: str) -> Dict[str, Any]:
    row = db.one("SELECT id FROM submissions WHERE id=?", (submission_id,))
    if not row:
        raise HTTPException(status_code=404, detail="提交记录不存在")
    refresh_submission_counts(submission_id)
    db.execute("UPDATE submissions SET status=?, submitted_at=? WHERE id=?", ("submitted", now_iso(), submission_id))
    return get_submission_result(submission_id)


@app.get("/api/submissions/{submission_id}/result")
def get_submission_result(submission_id: str) -> Dict[str, Any]:
    row = db.one("SELECT * FROM submissions WHERE id=?", (submission_id,))
    if not row:
        raise HTTPException(status_code=404, detail="提交记录不存在")
    answers = db.query("SELECT * FROM submission_answers WHERE submission_id=? ORDER BY created_at", (submission_id,))
    for answer in answers:
        answer["feedback"] = json_loads(answer.pop("feedback_json"), {})
        item = db.one("SELECT id,item_type,prompt_text,answer,audio_text,order_no,answer_meta_json FROM dictation_items WHERE id=?", (answer["item_id"],))
        if item:
            item["answer_meta"] = json_loads(item.pop("answer_meta_json"), {})
        answer["item"] = item
    row["answers"] = answers
    return row


@app.get("/api/student/mistakes")
def list_mistakes(student_id: str = Query(...)) -> List[Dict[str, Any]]:
    rows = db.query(
        "SELECT * FROM mistake_records WHERE student_id=? ORDER BY wrong_count DESC,last_wrong_at DESC",
        (student_id,),
    )
    for row in rows:
        source_item_ids = json_loads(row.pop("source_item_ids_json"), [])
        row["source_item_ids"] = source_item_ids
        row["first_wrong_at_display"] = format_beijing_time(row.get("first_wrong_at"))
        row["last_wrong_at_display"] = format_beijing_time(row.get("last_wrong_at"))
        row["next_review_at_display"] = format_beijing_time(row.get("next_review_at"))
        row["lesson"] = None
        row["source_task"] = None
        if source_item_ids:
            source = db.one(
                """
                SELECT dt.id AS task_id, dt.title AS task_title, lp.content_json
                FROM dictation_items di
                JOIN dictation_tasks dt ON di.task_id = dt.id
                JOIN learning_packs lp ON dt.learning_pack_id = lp.id
                WHERE di.id=?
                """,
                (source_item_ids[0],),
            )
            if source:
                content = json_loads(source.pop("content_json"), {})
                row["lesson"] = content.get("lesson")
                row["source_task"] = {"id": source["task_id"], "title": source["task_title"]}
    return rows


@app.patch("/api/student/mistakes/{mistake_id}")
def update_mistake_status(mistake_id: str, request: UpdateMistakeStatusRequest) -> Dict[str, Any]:
    mistake = db.one("SELECT * FROM mistake_records WHERE id=? AND student_id=?", (mistake_id, request.student_id))
    if not mistake:
        raise HTTPException(status_code=404, detail="错题不存在")
    next_review_at = mistake.get("next_review_at")
    if request.status == "corrected":
        next_review_at = (datetime.utcnow() + timedelta(days=1)).isoformat(timespec="seconds") + "Z"
    elif request.status == "reviewing":
        next_review_at = (datetime.utcnow() + timedelta(days=2)).isoformat(timespec="seconds") + "Z"
    elif request.status == "passed":
        next_review_at = None
    elif request.status == "repeated":
        next_review_at = (datetime.utcnow() + timedelta(hours=12)).isoformat(timespec="seconds") + "Z"
    db.execute(
        "UPDATE mistake_records SET status=?, next_review_at=? WHERE id=?",
        (request.status, next_review_at, mistake_id),
    )
    refreshed = [item for item in list_mistakes(request.student_id) if item["id"] == mistake_id]
    return refreshed[0] if refreshed else {"id": mistake_id, "status": request.status, "next_review_at": next_review_at}


@app.post("/api/student/mistakes/{mistake_id}/review")
def generate_review_exercises(mistake_id: str, request: ReviewRequest) -> Dict[str, Any]:
    mistake = db.one("SELECT * FROM mistake_records WHERE id=? AND student_id=?", (mistake_id, request.student_id))
    if not mistake:
        raise HTTPException(status_code=404, detail="错题不存在")
    value = mistake["char_or_word"]
    char = value[0]
    data = KNOWN_CHARS.get(char, {})
    confusing = [item["char"] for item in data.get("confusing_chars", [])]
    exercises = [
        {"type": "dictation", "prompt": f"再听写一次：{value}", "answer": value},
        {"type": "word", "prompt": f"请用“{char}”组一个词", "answer": "、".join(data.get("words", [value]))},
    ]
    if confusing:
        exercises.append(
            {
                "type": "choice",
                "prompt": f"下面哪个字适合“{data.get('simple_sentence', value)}”？",
                "options": [char] + confusing[:3],
                "answer": char,
            }
        )
    return {"mistake": mistake, "exercises": exercises}


def chinese_chars(value: str) -> List[str]:
    return [char for char in value if re.match(r"[\u4e00-\u9fff]", char)]


@app.get("/api/reports/tasks/{task_id}")
def get_task_report(task_id: str) -> Dict[str, Any]:
    task = get_dictation_task(task_id)
    pack = get_learning_pack(task["learning_pack_id"])
    lesson_info = pack.get("content", {}).get("lesson", {})
    students = list_students(task["class_id"])
    submissions = db.query("SELECT * FROM submissions WHERE task_id=? ORDER BY submitted_at DESC", (task_id,))
    answers = db.query(
        """
        SELECT sa.*, di.answer, di.prompt_text, di.item_type, s.student_id, u.name AS student_name
        FROM submission_answers sa
        JOIN dictation_items di ON sa.item_id = di.id
        JOIN submissions s ON sa.submission_id = s.id
        JOIN users u ON s.student_id = u.id
        WHERE s.task_id=?
        """,
        (task_id,),
    )
    total_students = len(students)
    submitted_count = len({row["student_id"] for row in submissions if row["status"] in {"submitted", "reviewed"}})
    total_answers = len(answers)
    correct_answers = sum(1 for row in answers if row["result"] == "correct")
    wrong_answers = [row for row in answers if row["result"] == "wrong"]
    mistake_rank: Dict[str, int] = {}
    char_rank: Dict[str, int] = {}
    word_rank: Dict[str, int] = {}
    mistake_types: Dict[str, int] = {}
    for row in wrong_answers:
        answer = row["answer"]
        mistake_rank[answer] = mistake_rank.get(answer, 0) + 1
        chars = chinese_chars(answer)
        for char in chars:
            char_rank[char] = char_rank.get(char, 0) + 1
        if len(chars) > 1:
            word_rank[answer] = word_rank.get(answer, 0) + 1
        mistake_type = row["mistake_type"] or "unknown"
        mistake_types[mistake_type] = mistake_types.get(mistake_type, 0) + 1
    top_mistakes = [{"item": key, "count": value} for key, value in sorted(mistake_rank.items(), key=lambda kv: kv[1], reverse=True)]
    top_chars = [{"item": key, "count": value} for key, value in sorted(char_rank.items(), key=lambda kv: kv[1], reverse=True)[:10]]
    top_words = [{"item": key, "count": value} for key, value in sorted(word_rank.items(), key=lambda kv: kv[1], reverse=True)[:10]]

    unit_bucket: Dict[str, Dict[str, Any]] = {}
    for item in task["items"]:
        unit_bucket[item["id"]] = {
            "item_id": item["id"],
            "order_no": item["order_no"],
            "item": item["answer"],
            "question_type": item["item_type"],
            "difficulty": item["difficulty"],
            "answered_count": 0,
            "correct_count": 0,
            "wrong_count": 0,
            "pending_review_count": 0,
        }
    for row in answers:
        bucket = unit_bucket.get(row["item_id"])
        if not bucket:
            continue
        bucket["answered_count"] += 1
        if row["result"] == "correct":
            bucket["correct_count"] += 1
        elif row["result"] == "pending_review":
            bucket["pending_review_count"] += 1
        else:
            bucket["wrong_count"] += 1
    unit_mastery = []
    for bucket in sorted(unit_bucket.values(), key=lambda item: item["order_no"]):
        answered_count = bucket["answered_count"]
        unit_mastery.append(
            {
                **bucket,
                "accuracy": round((bucket["correct_count"] / answered_count) * 100, 1) if answered_count else 0,
            }
        )

    def rank_for(mistake_type: str) -> List[Dict[str, Any]]:
        counter: Dict[str, int] = {}
        for row in wrong_answers:
            if row["mistake_type"] == mistake_type:
                counter[row["answer"]] = counter.get(row["answer"], 0) + 1
        return [{"item": key, "count": value} for key, value in sorted(counter.items(), key=lambda kv: kv[1], reverse=True)[:10]]

    weak_by_student: Dict[str, Dict[str, Dict[str, Any]]] = {}
    for row in wrong_answers:
        student_bucket = weak_by_student.setdefault(row["student_id"], {})
        item_bucket = student_bucket.setdefault(row["answer"], {"item": row["answer"], "count": 0, "types": set()})
        item_bucket["count"] += 1
        item_bucket["types"].add(row["mistake_type"] or "unknown")

    def weak_words_for(student_id: str) -> List[Dict[str, Any]]:
        bucket = weak_by_student.get(student_id, {})
        rows = []
        for item in bucket.values():
            rows.append({"item": item["item"], "count": item["count"], "types": sorted(item["types"])})
        return sorted(rows, key=lambda row: row["count"], reverse=True)[:5]

    suggestions = build_report_suggestions(top_mistakes, mistake_types)
    student_rows = []
    for student in students:
        sub = next((item for item in submissions if item["student_id"] == student["id"]), None)
        student_rows.append(
            {
                "student_id": student["id"],
                "student_name": student["name"],
                "status": sub["status"] if sub else "not_started",
                "correct_count": sub["correct_count"] if sub else 0,
                "total_count": sub["total_count"] if sub else len(task["items"]),
                "wrong_count": sub["wrong_count"] if sub else 0,
                "pending_review_count": sub["pending_review_count"] if sub else 0,
                "submitted_at": sub["submitted_at"] if sub else None,
                "accuracy": round((sub["correct_count"] / sub["total_count"]) * 100, 1) if sub and sub["total_count"] else 0,
                "weak_words": weak_words_for(student["id"]),
            }
        )
    return {
        "task": task,
        "lesson": lesson_info,
        "summary": {
            "total_students": total_students,
            "submitted_count": submitted_count,
            "completion_rate": round((submitted_count / total_students) * 100, 1) if total_students else 0,
            "total_answers": total_answers,
            "correct_answers": correct_answers,
            "accuracy": round((correct_answers / total_answers) * 100, 1) if total_answers else 0,
            "wrong_count": len(wrong_answers),
            "pending_review_count": sum(1 for row in answers if row["result"] == "pending_review"),
        },
        "top_mistakes": top_mistakes[:10],
        "top_chars": top_chars,
        "top_words": top_words,
        "unit_mastery": unit_mastery,
        "mistake_types": [{"type": key, "count": value} for key, value in sorted(mistake_types.items(), key=lambda kv: kv[1], reverse=True)],
        "similar_shape_rank": rank_for("similar_shape_confusion"),
        "homophone_rank": rank_for("homophone_confusion"),
        "students": student_rows,
        "suggestions": suggestions,
    }


def build_report_suggestions(top_mistakes: List[Dict[str, Any]], mistake_types: Dict[str, int]) -> List[Dict[str, str]]:
    suggestions: List[Dict[str, str]] = []
    if top_mistakes:
        focus = "、".join(item["item"] for item in top_mistakes[:3])
        suggestions.append(
            {
                "target": "全班",
                "focus": focus,
                "duration": "5-8 分钟",
                "activity": "错字重听写 + 组词复练",
                "reason": "这些字词在本次听写中出错次数较高。",
            }
        )
    if mistake_types.get("similar_shape_confusion"):
        suggestions.append(
            {
                "target": "易混淆学生",
                "focus": "形近字辨析",
                "duration": "5 分钟",
                "activity": "偏旁归类 + 语境选字",
                "reason": "本次存在形近字混淆，需要用偏旁和意思建立区分。",
            }
        )
    if mistake_types.get("homophone_confusion"):
        suggestions.append(
            {
                "target": "易混淆学生",
                "focus": "同音字辨析",
                "duration": "5 分钟",
                "activity": "听音选字 + 句子填空",
                "reason": "本次存在同音字混淆，需要结合语境练习。",
            }
        )
    if not suggestions:
        suggestions.append(
            {
                "target": "全班",
                "focus": "保持复习节奏",
                "duration": "3 分钟",
                "activity": "随机抽读 + 轻量复测",
                "reason": "当前任务错题较少，可通过短时复测巩固。",
            }
        )
    return suggestions


def submission_status_label(status: str) -> str:
    return {
        "submitted": "已提交",
        "in_progress": "进行中",
        "reviewed": "已复核",
        "not_started": "未开始",
    }.get(status, status)


def mistake_type_label(status: Optional[str]) -> str:
    return {
        "glyph_error": "字形错误",
        "homophone_confusion": "同音字混淆",
        "similar_shape_confusion": "形近字混淆",
        "radical_confusion": "偏旁混淆",
        "unknown": "待分析",
    }.get(status or "unknown", status or "待分析")


EXPORT_HEADER = ["班级", "学生姓名或匿名编号", "任务名称", "课文或单元", "完成时间", "总题数", "正确题数", "正确率", "错字列表", "高频错误类型", "是否完成订正"]


def build_export_data(task_id: str, anonymous: bool = False) -> Dict[str, Any]:
    report = get_task_report(task_id)
    class_row = db.one("SELECT name FROM classes WHERE id=?", (report["task"]["class_id"],))
    class_name = class_row["name"] if class_row else report["task"]["class_id"]
    lesson = report.get("lesson") or {}
    lesson_or_unit = " ".join(part for part in [lesson.get("unit"), lesson.get("title")] if part) or report["task"]["title"]
    rows: List[List[Any]] = []
    for index, row in enumerate(report["students"], start=1):
        submission = db.one(
            "SELECT * FROM submissions WHERE task_id=? AND student_id=? ORDER BY created_at DESC",
            (task_id, row["student_id"]),
        )
        answers = []
        if submission:
            answers = db.query(
                """
                SELECT sa.result, sa.mistake_type, di.answer
                FROM submission_answers sa
                JOIN dictation_items di ON sa.item_id = di.id
                WHERE sa.submission_id=?
                """,
                (submission["id"],),
            )
        wrong_answers = [answer for answer in answers if answer["result"] == "wrong"]
        wrong_words = [answer["answer"] for answer in wrong_answers]
        type_counter: Dict[str, int] = {}
        for answer in wrong_answers:
            label = mistake_type_label(answer["mistake_type"])
            type_counter[label] = type_counter.get(label, 0) + 1
        high_freq_types = "、".join(f"{key}({value})" for key, value in sorted(type_counter.items(), key=lambda kv: kv[1], reverse=True))
        if not wrong_words:
            correction_status = "无需订正"
        else:
            placeholders = ",".join("?" for _ in wrong_words)
            records = db.query(
                f"SELECT status FROM mistake_records WHERE student_id=? AND char_or_word IN ({placeholders})",
                [row["student_id"], *wrong_words],
            )
            correction_status = "已完成" if records and all(item["status"] in {"corrected", "passed"} for item in records) else "未完成"
        rows.append(
            [
                class_name,
                f"学生{index:02d}" if anonymous else row["student_name"],
                report["task"]["title"],
                lesson_or_unit,
                format_beijing_time(submission["submitted_at"]) if submission else "",
                row["total_count"],
                row["correct_count"],
                f"{row['accuracy']}%",
                "、".join(dict.fromkeys(wrong_words)),
                high_freq_types,
                correction_status,
            ]
        )
    return {"report": report, "class_name": class_name, "lesson_or_unit": lesson_or_unit, "rows": rows}


@app.post("/api/reports/tasks/{task_id}/export")
def export_task_report(task_id: str, anonymous: bool = Query(False)) -> StreamingResponse:
    export_data = build_export_data(task_id, anonymous)
    report = export_data["report"]
    lesson_or_unit = export_data["lesson_or_unit"]
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(EXPORT_HEADER)
    writer.writerows(export_data["rows"])
    writer.writerow([])
    writer.writerow(["任务名称", report["task"]["title"]])
    writer.writerow(["课文或单元", lesson_or_unit])
    writer.writerow(["完成率", f"{report['summary']['completion_rate']}%"])
    writer.writerow(["平均正确率", f"{report['summary']['accuracy']}%"])
    writer.writerow(["待复核数", report["summary"].get("pending_review_count", 0)])
    writer.writerow([])
    writer.writerow(["学生", "状态", "正确数", "总题数", "正确率"])
    for row in report["students"]:
        writer.writerow([row["student_name"], submission_status_label(row["status"]), row["correct_count"], row["total_count"], f"{row['accuracy']}%"])
    writer.writerow([])
    writer.writerow(["易错字词", "出错次数"])
    for row in report["top_mistakes"]:
        writer.writerow([row["item"], row["count"]])
    writer.writerow([])
    writer.writerow(["Top 易错字", "出错次数"])
    for row in report.get("top_chars", []):
        writer.writerow([row["item"], row["count"]])
    writer.writerow([])
    writer.writerow(["Top 易错词", "出错次数"])
    for row in report.get("top_words", []):
        writer.writerow([row["item"], row["count"]])
    writer.writerow([])
    writer.writerow(["单元掌握概览", "题型", "已答人数", "正确数", "错误数", "待复核数", "正确率"])
    for row in report.get("unit_mastery", []):
        writer.writerow(
            [
                row["item"],
                row["question_type"],
                row["answered_count"],
                row["correct_count"],
                row["wrong_count"],
                row["pending_review_count"],
                f"{row['accuracy']}%",
            ]
        )
    writer.writerow([])
    writer.writerow(["错误类型", "次数"])
    for row in report["mistake_types"]:
        writer.writerow([mistake_type_label(row["type"]), row["count"]])
    writer.writerow([])
    writer.writerow(["复习建议", "对象", "重点", "时长", "活动", "原因"])
    for index, item in enumerate(report["suggestions"], start=1):
        writer.writerow([index, item["target"], item["focus"], item["duration"], item["activity"], item["reason"]])
    output.seek(0)
    filename = f"dictation-report-{task_id}.csv"
    return StreamingResponse(
        iter([output.getvalue().encode("utf-8-sig")]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.post("/api/reports/tasks/{task_id}/export.xlsx")
def export_task_report_xlsx(task_id: str, anonymous: bool = Query(False)) -> StreamingResponse:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise HTTPException(status_code=500, detail="未安装 openpyxl，无法导出 Excel") from exc

    export_data = build_export_data(task_id, anonymous)
    report = export_data["report"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "听写记录"
    sheet.append(EXPORT_HEADER)
    for row in export_data["rows"]:
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="EAF4F8")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="17324D")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column_index, width in enumerate([16, 18, 28, 26, 20, 10, 10, 12, 28, 24, 16], start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.freeze_panes = "A2"

    summary = workbook.create_sheet("班级概览")
    summary_rows = [
        ["任务名称", report["task"]["title"]],
        ["课文或单元", export_data["lesson_or_unit"]],
        ["提交人数", f"{report['summary']['submitted_count']} / {report['summary']['total_students']}"],
        ["完成率", f"{report['summary']['completion_rate']}%"],
        ["平均正确率", f"{report['summary']['accuracy']}%"],
        ["错题数", report["summary"]["wrong_count"]],
        ["待复核数", report["summary"].get("pending_review_count", 0)],
    ]
    for row in summary_rows:
        summary.append(row)
    for cell in summary["A"]:
        cell.font = Font(bold=True, color="17324D")
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 42

    stats = workbook.create_sheet("易错统计")
    stats.append(["易错字词", "出错次数"])
    for row in report["top_mistakes"]:
        stats.append([row["item"], row["count"]])
    stats.append([])
    stats.append(["Top 易错字", "出错次数"])
    for row in report.get("top_chars", []):
        stats.append([row["item"], row["count"]])
    stats.append([])
    stats.append(["Top 易错词", "出错次数"])
    for row in report.get("top_words", []):
        stats.append([row["item"], row["count"]])
    stats.append([])
    stats.append(["单元掌握概览", "题型", "已答人数", "正确数", "错误数", "待复核数", "正确率"])
    for row in report.get("unit_mastery", []):
        stats.append(
            [
                row["item"],
                row["question_type"],
                row["answered_count"],
                row["correct_count"],
                row["wrong_count"],
                row["pending_review_count"],
                f"{row['accuracy']}%",
            ]
        )
    stats.append([])
    stats.append(["错误类型", "次数"])
    for row in report["mistake_types"]:
        stats.append([mistake_type_label(row["type"]), row["count"]])
    stats.append([])
    stats.append(["复习建议", "对象", "重点", "时长", "活动", "原因"])
    for index, item in enumerate(report["suggestions"], start=1):
        stats.append([index, item["target"], item["focus"], item["duration"], item["activity"], item["reason"]])
    for row in stats.iter_rows(min_row=1, max_row=stats.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column_index, width in enumerate([18, 14, 14, 12, 12, 12, 12], start=1):
        stats.column_dimensions[get_column_letter(column_index)].width = width

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"dictation-report-{task_id}.xlsx"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
